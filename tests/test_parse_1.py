"""test_parse_1.py —— 由 3 个单一行为文件合并

由 2026-10-21 的测试重组合并而来；每个来源文件的原首段说明保留为小节注释。
"""


from __future__ import annotations


import pytest
from dhole_mcp.parse import parse_file
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import re
import zipfile

pytestmark = pytest.mark.smoke


# ============================================================================
# 来源：test_a_docx_delivers_the_outline_its_file_declares.py
# ============================================================================
# 契约测试：`.docx` 交回的是**这份文件自己声明的**大纲层级。

docx = pytest.importorskip("docx", reason="the [all] extra provides python-docx")



def _docx(tmp_path, name: str, build) -> str:
    doc = Document()
    build(doc)
    path = tmp_path / name
    doc.save(str(path))
    return str(path)


def _body(text: str) -> list[str]:
    """The delivered content lines, receipt line dropped."""
    return [ln for ln in text.splitlines()[1:] if ln.strip()]


def _first(text: str) -> str:
    return text.splitlines()[0]


def _declare_outline(paragraph, val: int) -> None:
    """Write the level on the paragraph itself (no heading style involved)."""
    el = OxmlElement("w:outlineLvl")
    el.set(qn("w:val"), str(val))
    paragraph._p.get_or_add_pPr().append(el)


def _rename_style(doc, internal_name: str, new_name: str) -> None:
    """Change the display name a producer wrote into ``w:name``."""
    for st in doc.styles.element.findall(qn("w:style")):
        nm = st.find(qn("w:name"))
        if nm is not None and nm.get(qn("w:val")) == internal_name:
            nm.set(qn("w:val"), new_name)


def _strip_declared_levels(doc) -> None:
    """Delete every ``w:outlineLvl`` the styles declare (nothing left but names)."""
    for st in doc.styles.element.findall(qn("w:style")):
        ppr = st.find(qn("w:pPr"))
        if ppr is None:
            continue
        lvl = ppr.find(qn("w:outlineLvl"))
        while lvl is not None:
            ppr.remove(lvl)
            lvl = ppr.find(qn("w:outlineLvl"))


def _style_based_on(doc, style_id: str, ui_name: str, parent_id: str,
                    own_level: int | None = None) -> None:
    """A custom paragraph style that inherits its outline level from ``parent_id``."""
    st = OxmlElement("w:style")
    st.set(qn("w:type"), "paragraph")
    st.set(qn("w:styleId"), style_id)
    nm = OxmlElement("w:name")
    nm.set(qn("w:val"), ui_name)
    st.append(nm)
    based = OxmlElement("w:basedOn")
    based.set(qn("w:val"), parent_id)
    st.append(based)
    if own_level is not None:
        ppr = OxmlElement("w:pPr")
        lvl = OxmlElement("w:outlineLvl")
        lvl.set(qn("w:val"), str(own_level))
        ppr.append(lvl)
        st.append(ppr)
    doc.styles.element.append(st)


class TestTheFilesOwnOutlineLevels:
    """层级由文件的声明决定，不由样式名决定。"""

    def test_a_level_five_heading_keeps_its_level(self, tmp_path):
        """Word 的 1..9 级里我们从前只认 1..4：五级六级落成正文，读者看不出它是标题。"""
        path = _docx(tmp_path, "five.docx", lambda d: (
            d.add_heading("五级", level=5),
            d.add_paragraph("正文"),
        ))

        text, err = parse_file(path)

        assert err == ""
        assert _body(text) == ["##### 五级", "正文"]

    def test_a_level_six_heading_keeps_its_level(self, tmp_path):
        path = _docx(tmp_path, "six.docx", lambda d: (
            d.add_heading("一级", level=1),
            d.add_heading("六级", level=6),
            d.add_paragraph("正文"),
        ))

        text, _ = parse_file(path)

        assert _body(text) == ["# 一级", "###### 六级", "正文"]

    def test_a_heading_style_named_in_chinese_is_still_a_heading(self, tmp_path):
        """本地化改的是显示名，声明还在文件里 —— 从前名字一换，标题整行降级成正文。"""
        def build(d):
            d.add_heading("中文一级", level=1)
            d.add_paragraph("正文")
            d.add_heading("中文二级", level=2)
            _rename_style(d, "heading 1", "标题 1")
            _rename_style(d, "heading 2", "标题 2")

        text, _ = parse_file(_docx(tmp_path, "localized.docx", build))

        assert _body(text) == ["# 中文一级", "正文", "## 中文二级"]

    def test_a_style_built_on_heading_2_is_a_level_two_heading(self, tmp_path):
        """自定义样式（真稿子里「小标题」的常见做法）：层级从父样式继承下来。"""
        def build(d):
            _style_based_on(d, "MySection", "我的节标题", "Heading2")
            para = d.add_paragraph("继承来的二级")
            para.style = d.styles["我的节标题"]
            d.add_paragraph("正文")

        text, _ = parse_file(_docx(tmp_path, "inherited.docx", build))

        assert _body(text) == ["## 继承来的二级", "正文"]

    def test_a_paragraph_that_declares_its_level_without_a_heading_style(self, tmp_path):
        """有些导出器只往段落上写 `w:outlineLvl`，样式仍是 Normal。"""
        def build(d):
            para = d.add_paragraph("直接声明的二级")
            _declare_outline(para, 1)
            d.add_paragraph("正文")

        text, _ = parse_file(_docx(tmp_path, "direct.docx", build))

        assert _body(text) == ["## 直接声明的二级", "正文"]

    def test_the_declaration_beats_a_style_name_that_says_otherwise(self, tmp_path):
        """两个权威打架时听声明 —— 导航窗格也是这么显示的。"""
        def build(d):
            d.add_heading("冲突的一行", level=1)
            _rename_style(d, "heading 1", "heading 3")

        text, _ = parse_file(_docx(tmp_path, "conflict.docx", build))

        assert _body(text) == ["# 冲突的一行"]

    def test_a_paragraph_that_declares_itself_body_text_is_body_text(self, tmp_path):
        """`w:val="9"` 是 Word 拼写的"不进大纲"：段落这么说，就不给它横档。"""
        def build(d):
            para = d.add_heading("名字是一级，段落说自己不是", level=1)
            _declare_outline(para, 9)
            d.add_paragraph("正文")

        text, _ = parse_file(_docx(tmp_path, "lvl9.docx", build))

        assert _body(text) == ["名字是一级，段落说自己不是", "正文"]


class TestNothingThatReadsRightTodayIsLost:
    """修层级不许少标题：样式里查不到声明时，退回原来的名字映射。"""

    def test_headings_known_only_by_name_still_come_out_as_headings(self, tmp_path):
        """把样式里的 `w:outlineLvl` 全删光（残缺包 / 只写名字的导出器）：一级到四级照旧。"""
        def build(d):
            d.add_heading("一级", level=1)
            d.add_heading("二级", level=2)
            d.add_heading("三级", level=3)
            d.add_heading("四级", level=4)
            d.add_paragraph("正文")
            _strip_declared_levels(d)

        text, _ = parse_file(_docx(tmp_path, "names_only.docx", build))

        assert _body(text) == ["# 一级", "## 二级", "### 三级", "#### 四级", "正文"]

    def test_a_document_that_needs_no_note_is_delivered_unchanged(self, tmp_path):
        """没 Title、没深层级的普通稿：首行与正文逐字照旧。"""
        def build(d):
            d.add_heading("一级", level=1)
            d.add_paragraph("正文")
            d.add_heading("二级", level=2)
            d.add_paragraph("列表项", style="List Paragraph")

        text, _ = parse_file(_docx(tmp_path, "plain.docx", build))

        assert _first(text) == (
            "[Word document; 4 paragraph(s), 0 table(s) read; headers/footers, "
            "footnotes, comments and images not included]")
        assert _body(text) == ["# 一级", "正文", "## 二级", "- 列表项"]

    def test_a_table_still_sits_between_its_paragraphs(self, tmp_path):
        def build(d):
            d.add_heading("一级", level=1)
            table = d.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "甲"
            table.cell(0, 1).text = "乙"
            d.add_heading("五级", level=5)

        text, _ = parse_file(_docx(tmp_path, "table.docx", build))

        order = [text.index(m) for m in ("# 一级", "| 甲 |", "##### 五级")]
        assert order == sorted(order)

    def test_a_document_with_no_title_says_nothing_about_one(self, tmp_path):
        def build(d):
            d.add_heading("一级", level=1)
            d.add_paragraph("正文")

        text, _ = parse_file(_docx(tmp_path, "no_title.docx", build))

        assert "document title" not in _first(text)


class TestTheTitleIsToldApartFromTheSections:
    """Title 不在这份文件的大纲里，但它和一级标题落在同一格 `#` —— 得说出来。"""

    def test_a_title_and_its_level_one_headings_are_not_left_equal(self, tmp_path):
        """台账那一发实测的形状：`# 跨文化传播 · 第一周导读` / `# 阅读安排` / `# 作业`。"""
        def build(d):
            d.add_paragraph("跨文化传播 · 第一周导读", style="Title")
            d.add_paragraph("这一周只解决一个问题。")
            d.add_heading("阅读安排", level=1)
            d.add_paragraph("先读 Hall 1976。")
            d.add_heading("作业", level=1)

        text, _ = parse_file(_docx(tmp_path, "guide.docx", build))

        assert [ln for ln in _body(text) if ln.startswith("#")] == [
            "# 跨文化传播 · 第一周导读", "# 阅读安排", "# 作业"]
        assert "document title" in _first(text)
        assert "2 level-1 heading(s)" in _first(text)

    def test_a_title_with_no_level_one_heading_needs_no_note(self, tmp_path):
        """Title 之上没有别的东西挤同一格时，`#` 自己就说清了。"""
        def build(d):
            d.add_paragraph("文档标题", style="Title")
            d.add_heading("二级小节", level=2)

        text, _ = parse_file(_docx(tmp_path, "title_only.docx", build))

        assert _body(text) == ["# 文档标题", "## 二级小节"]
        assert "document title" not in _first(text)

    def test_the_real_fixture_shape_names_the_collision(self, tmp_path):
        """首行那句要能被读的人用来分家：说的是同一格 `#`，不是另开一格。"""
        def build(d):
            d.add_paragraph("标题行", style="Title")
            d.add_heading("第一节", level=1)

        text, _ = parse_file(_docx(tmp_path, "collide.docx", build))

        assert "# rung" in _first(text)
        assert "1 level-1 heading(s)" in _first(text)


class TestLevelsMarkdownCannotCarry:
    """Word 声明到九级，markdown 只有六格横档：跨不过去的那几行不许静默降级成正文。"""

    def test_levels_below_markdowns_deepest_rung_are_clamped_and_named(self, tmp_path):
        def build(d):
            d.add_heading("七级", level=7)
            d.add_heading("九级", level=9)
            d.add_paragraph("正文")

        text, _ = parse_file(_docx(tmp_path, "deep.docx", build))

        assert _body(text) == ["###### 七级", "###### 九级", "正文"]
        assert "2 deep heading(s) flattened to ######" in _first(text)

    def test_a_document_without_deep_levels_does_not_mention_them(self, tmp_path):
        def build(d):
            d.add_heading("六级", level=6)

        text, _ = parse_file(_docx(tmp_path, "not_deep.docx", build))

        assert _body(text) == ["###### 六级"]
        assert "flattened" not in _first(text)


class TestABrokenStyleChainDoesNotStopTheRead:
    """`w:basedOn` 成环（坏包）：走链子的那一段要能自己收住，整份文件照读。"""

    def test_a_cycle_in_based_on_still_delivers_every_paragraph(self, tmp_path):
        def build(d):
            _style_based_on(d, "LoopA", "环甲", "LoopB")
            _style_based_on(d, "LoopB", "环乙", "LoopA")
            para = d.add_paragraph("环上的段落")
            para.style = d.styles["环甲"]
            d.add_paragraph("正文")

        text, err = parse_file(_docx(tmp_path, "cycle.docx", build))

        assert err == ""
        assert "环上的段落" in text and "正文" in text


# ============================================================================
# 来源：test_the_notes_come_with_the_slide_they_belong_to.py
# ============================================================================
# 契约测试：`.pptx` 的讲者备注，以及 `<a:fld>` 布局字段不是正文。

_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PKG_RELS = "http://schemas.openxmlformats.org/package/2006/relationships"


def _fld(value: str, kind: str = "slidenum") -> str:
    """A layout placeholder — the shape PowerPoint writes for page numbers."""
    return (f'<a:p><a:fld id="{{{kind}-0000-0000-0000-000000000001}}" type="{kind}">'
            f"<a:t>{value}</a:t></a:fld></a:p>")


def _slide_xml(body: str, *, pagenum: str | None = None) -> str:
    extra = ""
    if pagenum is not None:
        extra = (f'<p:sp><p:nvSpPr><p:nvPr><p:ph type="sldNum"/></p:nvPr></p:nvSpPr>'
                 f"<p:txBody>{_fld(pagenum)}</p:txBody></p:sp>")
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<p:sld xmlns:a="{_A}" xmlns:p="{_P}"><p:cSld><p:spTree>'
            f'<p:sp><p:txBody><a:p><a:r><a:t>{body}</a:t></a:r></a:p></p:txBody></p:sp>'
            f"{extra}</p:spTree></p:cSld></p:sld>")


def _notes_xml(paragraphs: list[str], *, pagenum: str | None = None) -> str:
    body = "".join(f"<a:p><a:r><a:t>{p}</a:t></a:r></a:p>" for p in paragraphs)
    num = (f'<p:sp><p:nvSpPr><p:nvPr><p:ph type="sldNum"/></p:nvPr></p:nvSpPr>'
           f"<p:txBody>{_fld(pagenum)}</p:txBody></p:sp>") if pagenum else ""
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<p:notes xmlns:a="{_A}" xmlns:p="{_P}"><p:cSld><p:spTree>'
            f'<p:sp><p:nvSpPr><p:nvPr><p:ph type="body"/></p:nvPr></p:nvSpPr>'
            f"<p:txBody>{body}</p:txBody></p:sp>{num}</p:spTree></p:cSld></p:notes>")


def _deck(tmp_path, name: str, slides: dict[str, str], *,
          notes: dict[str, list[str]] | None = None,
          page_numbers: dict[str, str] | None = None,
          play_order: list[str] | None = None,
          notes_pagenum: bool = True) -> str:
    """A minimal .pptx with optional per-slide notes, built by hand."""
    notes = notes or {}
    page_numbers = page_numbers or {}
    parts = sorted(slides, key=lambda n: int(re.search(r"(\d+)", n).group(1)))
    play_order = play_order or parts
    rid_of = {part: f"rId{i + 1}" for i, part in enumerate(parts)}

    path = tmp_path / name
    with zipfile.ZipFile(path, "w") as zf:
        listed = "".join(
            f'<p:sldId id="{256 + i}" r:id="{rid_of[part]}"/>'
            for i, part in enumerate(play_order))
        zf.writestr(
            "ppt/presentation.xml",
            f'<?xml version="1.0"?><p:presentation xmlns:p="{_P}" xmlns:r="{_R}">'
            f"<p:sldIdLst>{listed}</p:sldIdLst></p:presentation>")
        zf.writestr(
            "ppt/_rels/presentation.xml.rels",
            f'<Relationships xmlns="{_PKG_RELS}">'
            + "".join(f'<Relationship Id="{rid}" Type="{_R}/slide" '
                      f'Target="{part[len("ppt/"):]}"/>'
                      for part, rid in rid_of.items())
            + "</Relationships>")

        for i, part in enumerate(parts, 1):
            zf.writestr(part, _slide_xml(slides[part],
                                         pagenum=page_numbers.get(part)))
            if part not in notes:
                continue
            note_part = f"ppt/notesSlides/notesSlide{i}.xml"
            zf.writestr(note_part, _notes_xml(
                notes[part], pagenum=str(i) if notes_pagenum else None))
            zf.writestr(
                f"{part.rsplit('/', 1)[0]}/_rels/{part.rsplit('/', 1)[-1]}.rels",
                f'<Relationships xmlns="{_PKG_RELS}"><Relationship Id="rId9" '
                f'Type="{_R}/notesSlide" Target="../notesSlides/'
                f'{note_part.rsplit("/", 1)[-1]}"/></Relationships>')
    return str(path)


def _blocks(text: str) -> list[list[str]]:
    """每一节从自己的**正文第一行**起算，装着它自己的正文与备注。

    窗口从前是部件名（`## slide1.xml`）—— 那既能当节的地址也能当顺序的读数。
    小节标题改成「第几页」之后（R02-025 那一把），部件名不再出场，这里改按
    **每页自己的文字**取节：它同样指得出是哪一页，备注挂在哪儿这条判据一点没变松。
    """
    out: list[list[str]] = []
    for ln in text.splitlines():
        if ln.startswith("## "):
            out.append([])
        elif out and ln.strip():
            out[-1].append(ln.strip())
    return out


def _by_body(text: str) -> dict:
    return {b[0]: b for b in _blocks(text) if b}


class TestNotesComeWithTheirSlide:
    """备注读进来，挂在它自己那一页下面，并且标清楚它是备注。"""

    def test_a_notes_part_is_emitted_under_its_own_slide(self, tmp_path):
        path = _deck(tmp_path, "notes.pptx",
                     {"ppt/slides/slide1.xml": "第三季度营收",
                      "ppt/slides/slide2.xml": "成本结构"},
                     notes={"ppt/slides/slide1.xml": ["这里要强调华东区的增长"]})

        text, err = parse_file(path)

        assert err == ""
        blocks = _by_body(text)
        assert blocks["第三季度营收"] == [
            "第三季度营收", "> Notes: 这里要强调华东区的增长"]
        # 另一页没有备注，它的块里就只有正文
        assert blocks["成本结构"] == ["成本结构"]

    def test_notes_are_labeled_so_they_are_not_mistaken_for_slide_text(self, tmp_path):
        """不标就和正文糊在一起 —— 读的人会把讲者备注当幻灯片内容引用。"""
        path = _deck(tmp_path, "labeled.pptx",
                     {"ppt/slides/slide1.xml": "第三季度营收"},
                     notes={"ppt/slides/slide1.xml": ["讲者备注一句"]})

        text, _ = parse_file(path)

        assert "> Notes: 讲者备注一句" in text

    def test_multi_paragraph_notes_keep_their_line_breaks(self, tmp_path):
        path = _deck(tmp_path, "multi.pptx",
                     {"ppt/slides/slide1.xml": "标题"},
                     notes={"ppt/slides/slide1.xml":
                            ["要点一", "数据来源：Q3 财报第 12 页"]})

        text, _ = parse_file(path)

        assert "> Notes: 要点一" in text
        assert "> 数据来源：Q3 财报第 12 页" in text

    def test_a_slide_without_notes_gets_no_placeholder(self, tmp_path):
        """同 `.docx` 的规矩：只在文件真有的时候才说。"""
        path = _deck(tmp_path, "mixed.pptx",
                     {"ppt/slides/slide1.xml": "有备注",
                      "ppt/slides/slide2.xml": "没备注"},
                     notes={"ppt/slides/slide1.xml": ["一句"]})

        text, _ = parse_file(path)

        assert "> Notes:" in text
        assert text.count("> Notes:") == 1, "没备注的页不该写占位符"
        assert "(no notes)" not in text

    def test_notes_follow_their_slide_when_the_deck_is_reordered(self, tmp_path):
        """备注跟的是**配对**，不是位置 —— 页序变了它得跟着自己那一页走。"""
        path = _deck(tmp_path, "reordered.pptx",
                     {"ppt/slides/slide1.xml": "甲页",
                      "ppt/slides/slide2.xml": "乙页"},
                     notes={"ppt/slides/slide2.xml": ["乙页的备注"]},
                     play_order=["ppt/slides/slide2.xml", "ppt/slides/slide1.xml"])

        text, _ = parse_file(path)

        assert [b[0] for b in _blocks(text)] == ["乙页", "甲页"], "页序没生效"
        blocks = _by_body(text)
        assert blocks["乙页"] == ["乙页", "> Notes: 乙页的备注"]
        assert blocks["甲页"] == ["甲页"]

    def test_a_notes_part_with_no_authored_text_gets_no_block(self, tmp_path):
        """备注页存在但里面只有编号字段：那不是备注，不该凭空长出一个 `> Notes:`。"""
        path = _deck(tmp_path, "empty-notes.pptx",
                     {"ppt/slides/slide1.xml": "正文"},
                     notes={"ppt/slides/slide1.xml": []})

        text, err = parse_file(path)

        assert err == ""
        assert "> Notes:" not in text

    def test_a_dangling_notes_relationship_is_skipped(self, tmp_path):
        """rels 指向一个不存在的备注部件：跳过它，别让整份文件读不出来。"""
        path = tmp_path / "dangling-notes.pptx"
        with zipfile.ZipFile(path, "w") as zf:
            zf.writestr("ppt/slides/slide1.xml", _slide_xml("正文"))
            zf.writestr(
                "ppt/slides/_rels/slide1.xml.rels",
                f'<Relationships xmlns="{_PKG_RELS}"><Relationship Id="rId9" '
                f'Type="{_R}/notesSlide" Target="../notesSlides/gone.xml"/>'
                "</Relationships>")

        text, err = parse_file(str(path))

        assert err == ""
        assert "正文" in text
        assert "> Notes:" not in text


class TestALayoutFieldIsNotText:
    """`<a:fld>` 是版式占位，不是作者写的字。"""

    def test_a_slide_number_placeholder_is_not_emitted(self, tmp_path):
        path = _deck(tmp_path, "pagenum.pptx",
                     {"ppt/slides/slide1.xml": "第三季度营收"},
                     page_numbers={"ppt/slides/slide1.xml": "3"})

        text, err = parse_file(path)

        assert err == ""
        # 旧行为：正文是 "第三季度营收" + 一个凭空多出来的 "3"
        assert "3" not in text, f"页码占位符漏进了正文：{text!r}"
        assert "第三季度营收" in text

    def test_a_notes_slide_number_placeholder_is_not_emitted(self, tmp_path):
        """备注页每一页都带编号字段 —— 不过滤它，读备注就变成读一堆数字。"""
        path = _deck(tmp_path, "notes-pagenum.pptx",
                     {"ppt/slides/slide1.xml": "正文"},
                     notes={"ppt/slides/slide1.xml": ["备注一句"]})

        text, err = parse_file(path)

        assert err == ""
        assert text.count("> Notes: 备注一句") == 1
        assert "\n> 1\n" not in text, "备注页的编号字段漏了出来"

    def test_a_date_field_is_dropped_too(self, tmp_path):
        """不止编号：日期占位符是同一类东西。"""
        path = tmp_path / "datefield.pptx"
        with zipfile.ZipFile(path, "w") as zf:
            zf.writestr("ppt/slides/slide1.xml",
                        '<?xml version="1.0"?>'
                        f'<p:sld xmlns:a="{_A}" xmlns:p="{_P}"><p:cSld><p:spTree>'
                        f'<p:sp><p:txBody><a:p><a:r><a:t>季度回顾</a:t></a:r></a:p>'
                        f"<a:p>{_fld('2026-10-06', 'datetime')}</a:p>"
                        f"</p:txBody></p:sp></p:spTree></p:cSld></p:sld>")

        text, err = parse_file(str(path))

        assert err == ""
        assert "季度回顾" in text
        assert "2026-10-06" not in text

    def test_text_beside_a_field_survives(self, tmp_path):
        """过滤的是字段本身，不是它所在的段落里的其它文字。"""
        path = tmp_path / "beside.pptx"
        with zipfile.ZipFile(path, "w") as zf:
            zf.writestr("ppt/slides/slide1.xml",
                        '<?xml version="1.0"?>'
                        f'<p:sld xmlns:a="{_A}" xmlns:p="{_P}"><p:cSld><p:spTree>'
                        f'<p:sp><p:txBody><a:p><a:r><a:t>第 </a:t></a:r>'
                        f'<a:fld id="{{X}}" type="slidenum"><a:t>7</a:t></a:fld>'
                        f'<a:r><a:t> 页结束</a:t></a:r></a:p>'
                        f"</p:txBody></p:sp></p:spTree></p:cSld></p:sld>")

        text, err = parse_file(str(path))

        assert err == ""
        assert "第" in text and "页结束" in text
        assert "7" not in text


class TestTheFirstLineSaysWhatIsLeft:
    """第一行的职责是声明没带什么 —— 备注带上了就从那张单子上下来。"""

    def test_the_first_line_no_longer_lists_speaker_notes(self, tmp_path):
        path = _deck(tmp_path, "first.pptx", {"ppt/slides/slide1.xml": "正文"})

        text, _ = parse_file(path)

        first = text.splitlines()[0]
        assert first.startswith("[1 slide(s);")
        assert "speaker notes" not in first

    def test_the_first_line_still_lists_what_really_is_dropped(self, tmp_path):
        """把备注从清单上划掉，不能顺手把版式也划掉 —— 那才是真的还在丢。"""
        path = _deck(tmp_path, "still.pptx", {"ppt/slides/slide1.xml": "正文"})

        text, _ = parse_file(path)

        first = text.splitlines()[0]
        assert "layout and diagram text are not included" in first


# ============================================================================
# 来源：test_a_docx_list_keeps_the_shape_the_file_declares.py
# ============================================================================
# 契约测试：`.docx` 的列表形状交回的是**这份文件声明的**形状。
# 重名消歧：TestNothingThatReadsRightTodayIsLost → TestNothingThatReadsRightTodayIsLost__a_docx_list_keeps_the_sh_b710、_body → _body__a_docx_list_keeps_the_sh_b710、docx → docx__a_docx_list_keeps_the_sh_b710

docx__a_docx_list_keeps_the_sh_b710 = pytest.importorskip("docx", reason="the [all] extra provides python-docx")



def _save(doc, tmp_path, name: str) -> str:
    path = tmp_path / name
    doc.save(str(path))
    return str(path)


def _body__a_docx_list_keeps_the_sh_b710(text: str) -> list[str]:
    return [ln for ln in text.splitlines()[1:] if ln.strip()]


def _declare_numpr(paragraph, ilvl: int | None, num_id: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    if ilvl is not None:
        lvl = OxmlElement("w:ilvl")
        lvl.set(qn("w:val"), str(ilvl))
        numpr.append(lvl)
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), num_id)
    numpr.append(nid)
    ppr.append(numpr)


def _style_with_numpr(doc, style_id: str, ui_name: str, num_id: str,
                      ilvl: int | None = None) -> None:
    """A custom paragraph style whose OWN pPr declares the numbering (no name clue)."""
    st = OxmlElement("w:style")
    st.set(qn("w:type"), "paragraph")
    st.set(qn("w:styleId"), style_id)
    nm = OxmlElement("w:name")
    nm.set(qn("w:val"), ui_name)
    st.append(nm)
    ppr = OxmlElement("w:pPr")
    numpr = OxmlElement("w:numPr")
    if ilvl is not None:
        lvl = OxmlElement("w:ilvl")
        lvl.set(qn("w:val"), str(ilvl))
        numpr.append(lvl)
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), num_id)
    numpr.append(nid)
    ppr.append(numpr)
    st.append(ppr)
    doc.styles.element.append(st)


def _rename_styles(doc, renames: dict[str, str]) -> None:
    for st in doc.styles.element.findall(qn("w:style")):
        nm = st.find(qn("w:name"))
        if nm is not None and nm.get(qn("w:val")) in renames:
            nm.set(qn("w:val"), renames[nm.get(qn("w:val"))])


def _rewrite_numbering(path: str, mutate) -> str:
    """Run `mutate` over the numbering part's XML and rebuild the package."""
    with zipfile.ZipFile(path) as zf:
        items = {n: zf.read(n) for n in zf.namelist()}
    part = next(n for n in items if n.endswith("numbering.xml"))
    items[part] = mutate(items[part].decode("utf-8")).encode("utf-8")
    out = path + ".mutated.docx"
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, blob in items.items():
            zf.writestr(name, blob)
    return out


class TestTheDeclaredListShape:
    """形状（圆点 / 编号 / 第几级）从字节里读，不从样式名里猜。"""

    def test_an_ordered_list_is_not_delivered_as_bullets(self, tmp_path):
        """`List Number` 的 numId 指向 `numFmt=decimal`：交回 `1.` 而不是 `·`。"""
        doc = Document()
        doc.add_paragraph("有序一", style="List Number")
        doc.add_paragraph("有序二", style="List Number")
        doc.add_paragraph("普通正文")

        text, err = parse_file(_save(doc, tmp_path, "number.docx"))

        assert err == ""
        assert _body__a_docx_list_keeps_the_sh_b710(text) == ["1. 有序一", "1. 有序二", "普通正文"]

    def test_an_unordered_list_still_comes_out_as_a_bullet(self, tmp_path):
        doc = Document()
        doc.add_paragraph("无序一", style="List Bullet")
        doc.add_paragraph("普通正文")

        text, _ = parse_file(_save(doc, tmp_path, "bullet.docx"))

        assert _body__a_docx_list_keeps_the_sh_b710(text) == ["- 无序一", "普通正文"]

    def test_a_localized_style_name_does_not_lose_the_marker(self, tmp_path):
        """改的是给人看的名，`w:numPr` 与 `numbering.xml` 一字未动。"""
        doc = Document()
        doc.add_paragraph("无序一", style="List Bullet")
        doc.add_paragraph("有序一", style="List Number")
        doc.add_paragraph("普通正文")
        _rename_styles(doc, {"List Bullet": "列表项目符号", "List Number": "列表编号",
                             "Normal": "正文"})

        text, _ = parse_file(_save(doc, tmp_path, "localized.docx"))

        assert _body__a_docx_list_keeps_the_sh_b710(text) == ["- 无序一", "1. 有序一", "普通正文"]

    def test_a_custom_style_that_only_declares_its_numbering(self, tmp_path):
        """自定义样式（名字里没有 list，声明在 styles.xml 的 pPr 上）。"""
        doc = Document()
        _style_with_numpr(doc, "MyItems", "我的条目", num_id="1")
        para = doc.add_paragraph("靠声明进列表的一")
        para.style = doc.styles["我的条目"]
        doc.add_paragraph("普通正文")

        text, _ = parse_file(_save(doc, tmp_path, "custom_style.docx"))

        assert _body__a_docx_list_keeps_the_sh_b710(text) == ["- 靠声明进列表的一", "普通正文"]

    def test_a_paragraph_that_declares_its_numbering_under_a_plain_style(self, tmp_path):
        doc = Document()
        para = doc.add_paragraph("段落自己声明的条目")
        _declare_numpr(para, ilvl=None, num_id="1")
        doc.add_paragraph("普通正文")

        text, _ = parse_file(_save(doc, tmp_path, "direct.docx"))

        assert _body__a_docx_list_keeps_the_sh_b710(text) == ["- 段落自己声明的条目", "普通正文"]

    def test_the_nesting_level_survives(self, tmp_path):
        """`w:ilvl` 是缩进层级：内层那一行要读得出它在外层下面。"""
        doc = Document()
        outer = doc.add_paragraph("外层条目")
        _declare_numpr(outer, ilvl=0, num_id="1")
        inner = doc.add_paragraph("内层条目")
        _declare_numpr(inner, ilvl=1, num_id="1")
        doc.add_paragraph("普通正文")

        text, _ = parse_file(_save(doc, tmp_path, "nested.docx"))

        assert _body__a_docx_list_keeps_the_sh_b710(text) == ["- 外层条目", "    - 内层条目", "普通正文"]

    def test_a_level_that_declares_no_marker_stays_body_text(self, tmp_path):
        """`w:numFmt="none"` 是 Word 拼写的"这一档不画记号"。"""
        doc = Document()
        para = doc.add_paragraph("落在 none 档的一行")
        _declare_numpr(para, ilvl=0, num_id="1")
        doc.add_paragraph("普通正文")
        path = _save(doc, tmp_path, "none_fmt.docx")

        def mutate(xml: str) -> str:
            return re.sub(
                r'(<w:abstractNum w:abstractNumId="8".*?<w:lvl w:ilvl="0"[^>]*>.*?)'
                r'<w:numFmt w:val="[^"]*"', r'\1<w:numFmt w:val="none"',
                xml, count=1, flags=re.S)

        text, _ = parse_file(_rewrite_numbering(path, mutate))

        assert _body__a_docx_list_keeps_the_sh_b710(text) == ["落在 none 档的一行", "普通正文"]


    def test_a_headed_line_wins_over_its_numbering(self, tmp_path):
        """既声明大纲层级又带着 numbering 的一行（Word 里就是"带编号的标题"）交标题那一格。

        这是一次**决定**而不是巧合：一句关于"这一行是第几级"的主张比"它前面画个记号"更强，
        交回的横档里已经容不下另一层身份。今天也是这个次序，这里把它钉住。
        """
        doc = Document()
        heading = doc.add_heading("第一章", level=1)
        _declare_numpr(heading, ilvl=0, num_id="1")
        doc.add_paragraph("普通正文")

        text, _ = parse_file(_save(doc, tmp_path, "numbered_heading.docx"))

        assert _body__a_docx_list_keeps_the_sh_b710(text) == ["# 第一章", "普通正文"]


class TestNothingThatReadsRightTodayIsLost__a_docx_list_keeps_the_sh_b710:
    """修形状不许少一行列表：读得出形状就交形状，读不出形状也仍说得出"这是一条列表"。"""

    def test_a_list_paragraph_with_no_declaration_keeps_the_name_fallback(self, tmp_path):
        """`List Paragraph` 什么都没说（样式与段落都无 numPr）：从前 `- `，往后也 `- `。"""
        doc = Document()
        doc.add_paragraph("列表项", style="List Paragraph")
        doc.add_paragraph("普通正文")

        text, _ = parse_file(_save(doc, tmp_path, "name_only.docx"))

        assert _body__a_docx_list_keeps_the_sh_b710(text) == ["- 列表项", "普通正文"]

    def test_a_package_without_numbering_still_says_it_is_a_list(self, tmp_path):
        """残缺包：`numPr` 在、`numbering.xml` 不在 ⇒ 形状读不出，但列表身份不许丢。"""
        doc = Document()
        doc.add_paragraph("无序一", style="List Bullet")
        doc.add_paragraph("普通正文")
        path = _save(doc, tmp_path, "no_numbering.docx")

        with zipfile.ZipFile(path) as zf:
            items = {n: zf.read(n) for n in zf.namelist()}
        drop = next(n for n in items if n.endswith("numbering.xml"))
        rels = next(n for n in items if n.endswith("_rels/document.xml.rels"))
        items[rels] = re.sub(r'<Relationship[^>]*numbering\.xml"[^>]*/>', "",
                             items[rels].decode("utf-8")).encode("utf-8")
        del items[drop]
        out = tmp_path / "stripped.docx"
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
            for name, blob in items.items():
                zf.writestr(name, blob)

        text, err = parse_file(str(out))

        assert err == ""
        assert _body__a_docx_list_keeps_the_sh_b710(text) == ["- 无序一", "普通正文"], "身份留着，形状不谎报"

    def test_prose_and_receipt_are_untouched(self, tmp_path):
        doc = Document()
        doc.add_heading("一级", level=1)
        doc.add_paragraph("有序一", style="List Number")
        doc.add_paragraph("正文一句。")

        text, err = parse_file(_save(doc, tmp_path, "plain.docx"))

        assert err == ""
        assert text.splitlines()[0].startswith("[Word document; 3 paragraph(s), 0 table(s) read")
        assert "# 一级" in text and "正文一句。" in text


class TestTheSameLockDoesNotMoveTheOtherShapes:
    """这一格只动列表：标题层级（R02-F44 那张脸）与表格顺序原样在。"""

    def test_headings_still_carry_the_declared_level(self, tmp_path):
        doc = Document()
        doc.add_heading("一级", level=1)
        doc.add_heading("五级", level=5)
        doc.add_paragraph("无序一", style="List Bullet")

        text, _ = parse_file(_save(doc, tmp_path, "heads.docx"))

        assert _body__a_docx_list_keeps_the_sh_b710(text) == ["# 一级", "##### 五级", "- 无序一"]

    def test_a_table_still_sits_where_the_document_put_it(self, tmp_path):
        doc = Document()
        doc.add_paragraph("无序一", style="List Bullet")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "甲"
        table.cell(0, 1).text = "乙"
        doc.add_paragraph("有序一", style="List Number")

        text, _ = parse_file(_save(doc, tmp_path, "table.docx"))

        order = [text.index(m) for m in ("- 无序一", "| 甲 |", "1. 有序一")]
        assert order == sorted(order)
