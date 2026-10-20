"""契约测试：`.docx` 的列表形状交回的是**这份文件声明的**形状。

`_parse_docx` 给列表加前缀靠一样东西：样式名里有没有 `"list"` 这个字面（`elif "list" in
style`）。而 OOXML 对"这是一条列表、第几级、圆点还是编号"有**权威字段**：段落或它所属样式上的
`w:numPr`（`w:ilvl` 是缩进层级、`w:numId` 指向 `word/numbering.xml` 里那份抽象编号定义，
那里的 `w:numFmt` 才说得出这是 `bullet` 还是 `decimal`）。样式名是给人看的显示串 ——
和 R02-F44 那一格同一个病：换个语言、换个自定义样式，声明还在字节里，读的人看不见了。

实测（本地合成稿，零网络，`review/.runtime/p4/f46_measure.py`）：

    `List Bullet`（样式带 numId=1 → bullet）      旧：`- 无序一`            形状靠名字蒙对
    `List Number`（样式带 numId=5 → **decimal**） 旧：`- 有序一`            ← 编号被交成圆点
    段落自己写 `numPr(ilvl=0/1)`                  旧：两行都是平铺的 `- `   ← 层级没了
    样式名改成「列表项目符号」/「列表编号」           旧：`无序一` / `有序一`    ← 前缀整行丢
    `List Paragraph` 而无任何 numPr               旧：`- 列表项`            那是名字蒙出来的

判据与 F44 同一条：**先问文件的声明，再问名字**。声明读得出形状就交形状（`1. ` vs `- `），
读不出形状（`numbering.xml` 不在、numId 指向空的）就交"这是一条列表"而不谎报它是编号；
两处都查不到声明时才退回原来的名字映射 —— 修形状不许让一行列表消失。
`w:numFmt="none"` 是 Word 拼写的"这一档不画记号"，段落落在这一档就是正文。
"""

from __future__ import annotations

import re
import zipfile

import pytest

from dhole_mcp.parse import parse_file

docx = pytest.importorskip("docx", reason="the [all] extra provides python-docx")

from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402


def _save(doc, tmp_path, name: str) -> str:
    path = tmp_path / name
    doc.save(str(path))
    return str(path)


def _body(text: str) -> list[str]:
    return [ln for ln in text.splitlines()[1:] if ln.strip()]


def _declare_numpr(paragraph, ilvl: int | None, num_id: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    if ilvl is not None:
        lvl = OxmlElement("w:ilvl")
        lvl.set(qn("w:val"), str(ilvl))
        numpr.append(lvl)
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), num_id)
    numpr.append(nid)
    ppr.append(numpr)


def _style_with_numpr(doc, style_id: str, ui_name: str, num_id: str,
                      ilvl: int | None = None) -> None:
    """A custom paragraph style whose OWN pPr declares the numbering (no name clue)."""
    st = OxmlElement("w:style")
    st.set(qn("w:type"), "paragraph")
    st.set(qn("w:styleId"), style_id)
    nm = OxmlElement("w:name")
    nm.set(qn("w:val"), ui_name)
    st.append(nm)
    ppr = OxmlElement("w:pPr")
    numpr = OxmlElement("w:numPr")
    if ilvl is not None:
        lvl = OxmlElement("w:ilvl")
        lvl.set(qn("w:val"), str(ilvl))
        numpr.append(lvl)
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), num_id)
    numpr.append(nid)
    ppr.append(numpr)
    st.append(ppr)
    doc.styles.element.append(st)


def _rename_styles(doc, renames: dict[str, str]) -> None:
    for st in doc.styles.element.findall(qn("w:style")):
        nm = st.find(qn("w:name"))
        if nm is not None and nm.get(qn("w:val")) in renames:
            nm.set(qn("w:val"), renames[nm.get(qn("w:val"))])


def _rewrite_numbering(path: str, mutate) -> str:
    """Run `mutate` over the numbering part's XML and rebuild the package."""
    with zipfile.ZipFile(path) as zf:
        items = {n: zf.read(n) for n in zf.namelist()}
    part = next(n for n in items if n.endswith("numbering.xml"))
    items[part] = mutate(items[part].decode("utf-8")).encode("utf-8")
    out = path + ".mutated.docx"
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, blob in items.items():
            zf.writestr(name, blob)
    return out


class TestTheDeclaredListShape:
    """形状（圆点 / 编号 / 第几级）从字节里读，不从样式名里猜。"""

    def test_an_ordered_list_is_not_delivered_as_bullets(self, tmp_path):
        """`List Number` 的 numId 指向 `numFmt=decimal`：交回 `1.` 而不是 `·`。"""
        doc = Document()
        doc.add_paragraph("有序一", style="List Number")
        doc.add_paragraph("有序二", style="List Number")
        doc.add_paragraph("普通正文")

        text, err = parse_file(_save(doc, tmp_path, "number.docx"))

        assert err == ""
        assert _body(text) == ["1. 有序一", "1. 有序二", "普通正文"]

    def test_an_unordered_list_still_comes_out_as_a_bullet(self, tmp_path):
        doc = Document()
        doc.add_paragraph("无序一", style="List Bullet")
        doc.add_paragraph("普通正文")

        text, _ = parse_file(_save(doc, tmp_path, "bullet.docx"))

        assert _body(text) == ["- 无序一", "普通正文"]

    def test_a_localized_style_name_does_not_lose_the_marker(self, tmp_path):
        """改的是给人看的名，`w:numPr` 与 `numbering.xml` 一字未动。"""
        doc = Document()
        doc.add_paragraph("无序一", style="List Bullet")
        doc.add_paragraph("有序一", style="List Number")
        doc.add_paragraph("普通正文")
        _rename_styles(doc, {"List Bullet": "列表项目符号", "List Number": "列表编号",
                             "Normal": "正文"})

        text, _ = parse_file(_save(doc, tmp_path, "localized.docx"))

        assert _body(text) == ["- 无序一", "1. 有序一", "普通正文"]

    def test_a_custom_style_that_only_declares_its_numbering(self, tmp_path):
        """自定义样式（名字里没有 list，声明在 styles.xml 的 pPr 上）。"""
        doc = Document()
        _style_with_numpr(doc, "MyItems", "我的条目", num_id="1")
        para = doc.add_paragraph("靠声明进列表的一")
        para.style = doc.styles["我的条目"]
        doc.add_paragraph("普通正文")

        text, _ = parse_file(_save(doc, tmp_path, "custom_style.docx"))

        assert _body(text) == ["- 靠声明进列表的一", "普通正文"]

    def test_a_paragraph_that_declares_its_numbering_under_a_plain_style(self, tmp_path):
        doc = Document()
        para = doc.add_paragraph("段落自己声明的条目")
        _declare_numpr(para, ilvl=None, num_id="1")
        doc.add_paragraph("普通正文")

        text, _ = parse_file(_save(doc, tmp_path, "direct.docx"))

        assert _body(text) == ["- 段落自己声明的条目", "普通正文"]

    def test_the_nesting_level_survives(self, tmp_path):
        """`w:ilvl` 是缩进层级：内层那一行要读得出它在外层下面。"""
        doc = Document()
        outer = doc.add_paragraph("外层条目")
        _declare_numpr(outer, ilvl=0, num_id="1")
        inner = doc.add_paragraph("内层条目")
        _declare_numpr(inner, ilvl=1, num_id="1")
        doc.add_paragraph("普通正文")

        text, _ = parse_file(_save(doc, tmp_path, "nested.docx"))

        assert _body(text) == ["- 外层条目", "    - 内层条目", "普通正文"]

    def test_a_level_that_declares_no_marker_stays_body_text(self, tmp_path):
        """`w:numFmt="none"` 是 Word 拼写的"这一档不画记号"。"""
        doc = Document()
        para = doc.add_paragraph("落在 none 档的一行")
        _declare_numpr(para, ilvl=0, num_id="1")
        doc.add_paragraph("普通正文")
        path = _save(doc, tmp_path, "none_fmt.docx")

        def mutate(xml: str) -> str:
            return re.sub(
                r'(<w:abstractNum w:abstractNumId="8".*?<w:lvl w:ilvl="0"[^>]*>.*?)'
                r'<w:numFmt w:val="[^"]*"', r'\1<w:numFmt w:val="none"',
                xml, count=1, flags=re.S)

        text, _ = parse_file(_rewrite_numbering(path, mutate))

        assert _body(text) == ["落在 none 档的一行", "普通正文"]


    def test_a_headed_line_wins_over_its_numbering(self, tmp_path):
        """既声明大纲层级又带着 numbering 的一行（Word 里就是"带编号的标题"）交标题那一格。

        这是一次**决定**而不是巧合：一句关于"这一行是第几级"的主张比"它前面画个记号"更强，
        交回的横档里已经容不下另一层身份。今天也是这个次序，这里把它钉住。
        """
        doc = Document()
        heading = doc.add_heading("第一章", level=1)
        _declare_numpr(heading, ilvl=0, num_id="1")
        doc.add_paragraph("普通正文")

        text, _ = parse_file(_save(doc, tmp_path, "numbered_heading.docx"))

        assert _body(text) == ["# 第一章", "普通正文"]


class TestNothingThatReadsRightTodayIsLost:
    """修形状不许少一行列表：读得出形状就交形状，读不出形状也仍说得出"这是一条列表"。"""

    def test_a_list_paragraph_with_no_declaration_keeps_the_name_fallback(self, tmp_path):
        """`List Paragraph` 什么都没说（样式与段落都无 numPr）：从前 `- `，往后也 `- `。"""
        doc = Document()
        doc.add_paragraph("列表项", style="List Paragraph")
        doc.add_paragraph("普通正文")

        text, _ = parse_file(_save(doc, tmp_path, "name_only.docx"))

        assert _body(text) == ["- 列表项", "普通正文"]

    def test_a_package_without_numbering_still_says_it_is_a_list(self, tmp_path):
        """残缺包：`numPr` 在、`numbering.xml` 不在 ⇒ 形状读不出，但列表身份不许丢。"""
        doc = Document()
        doc.add_paragraph("无序一", style="List Bullet")
        doc.add_paragraph("普通正文")
        path = _save(doc, tmp_path, "no_numbering.docx")

        with zipfile.ZipFile(path) as zf:
            items = {n: zf.read(n) for n in zf.namelist()}
        drop = next(n for n in items if n.endswith("numbering.xml"))
        rels = next(n for n in items if n.endswith("_rels/document.xml.rels"))
        items[rels] = re.sub(r'<Relationship[^>]*numbering\.xml"[^>]*/>', "",
                             items[rels].decode("utf-8")).encode("utf-8")
        del items[drop]
        out = tmp_path / "stripped.docx"
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
            for name, blob in items.items():
                zf.writestr(name, blob)

        text, err = parse_file(str(out))

        assert err == ""
        assert _body(text) == ["- 无序一", "普通正文"], "身份留着，形状不谎报"

    def test_prose_and_receipt_are_untouched(self, tmp_path):
        doc = Document()
        doc.add_heading("一级", level=1)
        doc.add_paragraph("有序一", style="List Number")
        doc.add_paragraph("正文一句。")

        text, err = parse_file(_save(doc, tmp_path, "plain.docx"))

        assert err == ""
        assert text.splitlines()[0].startswith("[Word document; 3 paragraph(s), 0 table(s) read")
        assert "# 一级" in text and "正文一句。" in text


class TestTheSameLockDoesNotMoveTheOtherShapes:
    """这一格只动列表：标题层级（R02-F44 那张脸）与表格顺序原样在。"""

    def test_headings_still_carry_the_declared_level(self, tmp_path):
        doc = Document()
        doc.add_heading("一级", level=1)
        doc.add_heading("五级", level=5)
        doc.add_paragraph("无序一", style="List Bullet")

        text, _ = parse_file(_save(doc, tmp_path, "heads.docx"))

        assert _body(text) == ["# 一级", "##### 五级", "- 无序一"]

    def test_a_table_still_sits_where_the_document_put_it(self, tmp_path):
        doc = Document()
        doc.add_paragraph("无序一", style="List Bullet")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "甲"
        table.cell(0, 1).text = "乙"
        doc.add_paragraph("有序一", style="List Number")

        text, _ = parse_file(_save(doc, tmp_path, "table.docx"))

        order = [text.index(m) for m in ("- 无序一", "| 甲 |", "1. 有序一")]
        assert order == sorted(order)
