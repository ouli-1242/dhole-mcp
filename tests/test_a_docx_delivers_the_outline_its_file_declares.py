"""契约测试：`.docx` 交回的是**这份文件自己声明的**大纲层级。

`_parse_docx` 从前依据**样式名**判层级：`"heading 1" in style.name`，只列了 1..4 四种。
样式名是给人看的显示串 —— 它会被本地化（中文 Word 写「标题 1」）、会被自定义样式换掉
（基于 Heading 2 的「我的节标题」）、也可能与文件声明冲突。而 OOXML 里"这一行是第几级"
有**权威字段**：`w:outlineLvl`（0 = 一级），写在段落样式上、或经 `w:basedOn` 从父样式继承、
或直接写在段落自己的 `w:pPr` 上；`w:val="9"` 是 Word 自己拼写的"这一行不进大纲"。
Word 的导航窗格读的就是这个字段。

实测（本地合成稿，零网络）：

    Heading 5 / 6            样式声明 outlineLvl 4 / 5      旧：正文两行（层级静默丢掉）
    Heading 7..9             样式声明 outlineLvl 6..8        旧：正文
    样式名改成「标题 1」        样式声明仍在                   旧：正文
    basedOn=Heading2 的自定义样式 继承来的声明                 旧：正文
    Normal + 段落直接 outlineLvl=1 段落自己的声明               旧：正文
    样式 id=Heading1 而名字写着 heading 3   两个权威打架        旧：###（名字赢）
    Title + 两个 Heading 1    Title 不在大纲里               旧：三行都是 #（谁是谁读不出）

判据落在**两件事**上，缺一不可：① 交回的层级来自文件的声明（名字与声明打架时听声明）；
② **今天靠名字认对的那一批一行都不许退回正文** —— 样式里查不到声明时退回原来的名字映射，
修"层级错"不能换来"标题少几行"。带不走的两件事（Title 挤在同一格 `#`、markdown 没有
第六格以下的横档）写在首行，不静默。
"""

from __future__ import annotations

import pytest

from dhole_mcp.parse import parse_file

docx = pytest.importorskip("docx", reason="the [all] extra provides python-docx")

from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402


def _docx(tmp_path, name: str, build) -> str:
    doc = Document()
    build(doc)
    path = tmp_path / name
    doc.save(str(path))
    return str(path)


def _body(text: str) -> list[str]:
    """The delivered content lines, receipt line dropped."""
    return [ln for ln in text.splitlines()[1:] if ln.strip()]


def _first(text: str) -> str:
    return text.splitlines()[0]


def _declare_outline(paragraph, val: int) -> None:
    """Write the level on the paragraph itself (no heading style involved)."""
    el = OxmlElement("w:outlineLvl")
    el.set(qn("w:val"), str(val))
    paragraph._p.get_or_add_pPr().append(el)


def _rename_style(doc, internal_name: str, new_name: str) -> None:
    """Change the display name a producer wrote into ``w:name``."""
    for st in doc.styles.element.findall(qn("w:style")):
        nm = st.find(qn("w:name"))
        if nm is not None and nm.get(qn("w:val")) == internal_name:
            nm.set(qn("w:val"), new_name)


def _strip_declared_levels(doc) -> None:
    """Delete every ``w:outlineLvl`` the styles declare (nothing left but names)."""
    for st in doc.styles.element.findall(qn("w:style")):
        ppr = st.find(qn("w:pPr"))
        if ppr is None:
            continue
        lvl = ppr.find(qn("w:outlineLvl"))
        while lvl is not None:
            ppr.remove(lvl)
            lvl = ppr.find(qn("w:outlineLvl"))


def _style_based_on(doc, style_id: str, ui_name: str, parent_id: str,
                    own_level: int | None = None) -> None:
    """A custom paragraph style that inherits its outline level from ``parent_id``."""
    st = OxmlElement("w:style")
    st.set(qn("w:type"), "paragraph")
    st.set(qn("w:styleId"), style_id)
    nm = OxmlElement("w:name")
    nm.set(qn("w:val"), ui_name)
    st.append(nm)
    based = OxmlElement("w:basedOn")
    based.set(qn("w:val"), parent_id)
    st.append(based)
    if own_level is not None:
        ppr = OxmlElement("w:pPr")
        lvl = OxmlElement("w:outlineLvl")
        lvl.set(qn("w:val"), str(own_level))
        ppr.append(lvl)
        st.append(ppr)
    doc.styles.element.append(st)


class TestTheFilesOwnOutlineLevels:
    """层级由文件的声明决定，不由样式名决定。"""

    def test_a_level_five_heading_keeps_its_level(self, tmp_path):
        """Word 的 1..9 级里我们从前只认 1..4：五级六级落成正文，读者看不出它是标题。"""
        path = _docx(tmp_path, "five.docx", lambda d: (
            d.add_heading("五级", level=5),
            d.add_paragraph("正文"),
        ))

        text, err = parse_file(path)

        assert err == ""
        assert _body(text) == ["##### 五级", "正文"]

    def test_a_level_six_heading_keeps_its_level(self, tmp_path):
        path = _docx(tmp_path, "six.docx", lambda d: (
            d.add_heading("一级", level=1),
            d.add_heading("六级", level=6),
            d.add_paragraph("正文"),
        ))

        text, _ = parse_file(path)

        assert _body(text) == ["# 一级", "###### 六级", "正文"]

    def test_a_heading_style_named_in_chinese_is_still_a_heading(self, tmp_path):
        """本地化改的是显示名，声明还在文件里 —— 从前名字一换，标题整行降级成正文。"""
        def build(d):
            d.add_heading("中文一级", level=1)
            d.add_paragraph("正文")
            d.add_heading("中文二级", level=2)
            _rename_style(d, "heading 1", "标题 1")
            _rename_style(d, "heading 2", "标题 2")

        text, _ = parse_file(_docx(tmp_path, "localized.docx", build))

        assert _body(text) == ["# 中文一级", "正文", "## 中文二级"]

    def test_a_style_built_on_heading_2_is_a_level_two_heading(self, tmp_path):
        """自定义样式（真稿子里「小标题」的常见做法）：层级从父样式继承下来。"""
        def build(d):
            _style_based_on(d, "MySection", "我的节标题", "Heading2")
            para = d.add_paragraph("继承来的二级")
            para.style = d.styles["我的节标题"]
            d.add_paragraph("正文")

        text, _ = parse_file(_docx(tmp_path, "inherited.docx", build))

        assert _body(text) == ["## 继承来的二级", "正文"]

    def test_a_paragraph_that_declares_its_level_without_a_heading_style(self, tmp_path):
        """有些导出器只往段落上写 `w:outlineLvl`，样式仍是 Normal。"""
        def build(d):
            para = d.add_paragraph("直接声明的二级")
            _declare_outline(para, 1)
            d.add_paragraph("正文")

        text, _ = parse_file(_docx(tmp_path, "direct.docx", build))

        assert _body(text) == ["## 直接声明的二级", "正文"]

    def test_the_declaration_beats_a_style_name_that_says_otherwise(self, tmp_path):
        """两个权威打架时听声明 —— 导航窗格也是这么显示的。"""
        def build(d):
            d.add_heading("冲突的一行", level=1)
            _rename_style(d, "heading 1", "heading 3")

        text, _ = parse_file(_docx(tmp_path, "conflict.docx", build))

        assert _body(text) == ["# 冲突的一行"]

    def test_a_paragraph_that_declares_itself_body_text_is_body_text(self, tmp_path):
        """`w:val="9"` 是 Word 拼写的"不进大纲"：段落这么说，就不给它横档。"""
        def build(d):
            para = d.add_heading("名字是一级，段落说自己不是", level=1)
            _declare_outline(para, 9)
            d.add_paragraph("正文")

        text, _ = parse_file(_docx(tmp_path, "lvl9.docx", build))

        assert _body(text) == ["名字是一级，段落说自己不是", "正文"]


class TestNothingThatReadsRightTodayIsLost:
    """修层级不许少标题：样式里查不到声明时，退回原来的名字映射。"""

    def test_headings_known_only_by_name_still_come_out_as_headings(self, tmp_path):
        """把样式里的 `w:outlineLvl` 全删光（残缺包 / 只写名字的导出器）：一级到四级照旧。"""
        def build(d):
            d.add_heading("一级", level=1)
            d.add_heading("二级", level=2)
            d.add_heading("三级", level=3)
            d.add_heading("四级", level=4)
            d.add_paragraph("正文")
            _strip_declared_levels(d)

        text, _ = parse_file(_docx(tmp_path, "names_only.docx", build))

        assert _body(text) == ["# 一级", "## 二级", "### 三级", "#### 四级", "正文"]

    def test_a_document_that_needs_no_note_is_delivered_unchanged(self, tmp_path):
        """没 Title、没深层级的普通稿：首行与正文逐字照旧。"""
        def build(d):
            d.add_heading("一级", level=1)
            d.add_paragraph("正文")
            d.add_heading("二级", level=2)
            d.add_paragraph("列表项", style="List Paragraph")

        text, _ = parse_file(_docx(tmp_path, "plain.docx", build))

        assert _first(text) == (
            "[Word document; 4 paragraph(s), 0 table(s) read; headers/footers, "
            "footnotes, comments and images not included]")
        assert _body(text) == ["# 一级", "正文", "## 二级", "- 列表项"]

    def test_a_table_still_sits_between_its_paragraphs(self, tmp_path):
        def build(d):
            d.add_heading("一级", level=1)
            table = d.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "甲"
            table.cell(0, 1).text = "乙"
            d.add_heading("五级", level=5)

        text, _ = parse_file(_docx(tmp_path, "table.docx", build))

        order = [text.index(m) for m in ("# 一级", "| 甲 |", "##### 五级")]
        assert order == sorted(order)

    def test_a_document_with_no_title_says_nothing_about_one(self, tmp_path):
        def build(d):
            d.add_heading("一级", level=1)
            d.add_paragraph("正文")

        text, _ = parse_file(_docx(tmp_path, "no_title.docx", build))

        assert "document title" not in _first(text)


class TestTheTitleIsToldApartFromTheSections:
    """Title 不在这份文件的大纲里，但它和一级标题落在同一格 `#` —— 得说出来。"""

    def test_a_title_and_its_level_one_headings_are_not_left_equal(self, tmp_path):
        """台账那一发实测的形状：`# 跨文化传播 · 第一周导读` / `# 阅读安排` / `# 作业`。"""
        def build(d):
            d.add_paragraph("跨文化传播 · 第一周导读", style="Title")
            d.add_paragraph("这一周只解决一个问题。")
            d.add_heading("阅读安排", level=1)
            d.add_paragraph("先读 Hall 1976。")
            d.add_heading("作业", level=1)

        text, _ = parse_file(_docx(tmp_path, "guide.docx", build))

        assert [ln for ln in _body(text) if ln.startswith("#")] == [
            "# 跨文化传播 · 第一周导读", "# 阅读安排", "# 作业"]
        assert "document title" in _first(text)
        assert "2 level-1 heading(s)" in _first(text)

    def test_a_title_with_no_level_one_heading_needs_no_note(self, tmp_path):
        """Title 之上没有别的东西挤同一格时，`#` 自己就说清了。"""
        def build(d):
            d.add_paragraph("文档标题", style="Title")
            d.add_heading("二级小节", level=2)

        text, _ = parse_file(_docx(tmp_path, "title_only.docx", build))

        assert _body(text) == ["# 文档标题", "## 二级小节"]
        assert "document title" not in _first(text)

    def test_the_real_fixture_shape_names_the_collision(self, tmp_path):
        """首行那句要能被读的人用来分家：说的是同一格 `#`，不是另开一格。"""
        def build(d):
            d.add_paragraph("标题行", style="Title")
            d.add_heading("第一节", level=1)

        text, _ = parse_file(_docx(tmp_path, "collide.docx", build))

        assert "# rung" in _first(text)
        assert "1 level-1 heading(s)" in _first(text)


class TestLevelsMarkdownCannotCarry:
    """Word 声明到九级，markdown 只有六格横档：跨不过去的那几行不许静默降级成正文。"""

    def test_levels_below_markdowns_deepest_rung_are_clamped_and_named(self, tmp_path):
        def build(d):
            d.add_heading("七级", level=7)
            d.add_heading("九级", level=9)
            d.add_paragraph("正文")

        text, _ = parse_file(_docx(tmp_path, "deep.docx", build))

        assert _body(text) == ["###### 七级", "###### 九级", "正文"]
        assert "2 deep heading(s) flattened to ######" in _first(text)

    def test_a_document_without_deep_levels_does_not_mention_them(self, tmp_path):
        def build(d):
            d.add_heading("六级", level=6)

        text, _ = parse_file(_docx(tmp_path, "not_deep.docx", build))

        assert _body(text) == ["###### 六级"]
        assert "flattened" not in _first(text)


class TestABrokenStyleChainDoesNotStopTheRead:
    """`w:basedOn` 成环（坏包）：走链子的那一段要能自己收住，整份文件照读。"""

    def test_a_cycle_in_based_on_still_delivers_every_paragraph(self, tmp_path):
        def build(d):
            _style_based_on(d, "LoopA", "环甲", "LoopB")
            _style_based_on(d, "LoopB", "环乙", "LoopA")
            para = d.add_paragraph("环上的段落")
            para.style = d.styles["环甲"]
            d.add_paragraph("正文")

        text, err = parse_file(_docx(tmp_path, "cycle.docx", build))

        assert err == ""
        assert "环上的段落" in text and "正文" in text
